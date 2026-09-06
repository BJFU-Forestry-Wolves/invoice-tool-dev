from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
from collections import defaultdict
from dataclasses import asdict, dataclass
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from order_date.file_grouper import find_files_by_bx_id
from order_date.input_loader import normalize_uploader
from rename_invoices_orders import count_comma_segments, render_target_filename, safe_text


PROJECT_ROOT = Path(__file__).resolve().parent
CLASSIFICATION_NAME_TEMPLATE = "{下单日期}_{编号}_{上传人}_{用途}_{金额}_{附件标记}"
LOW_VALUE_INVOICE_ONLY_LIMIT = 200.0
RESOLVED_REVIEW_STATUSES = {"确认无误", "已修正"}
SPECIAL_REVIEW_KEYWORDS = ("特殊命名标记", "理论上不能报销", "已经提交")
MANUAL_PREFIX_RE = re.compile(r"^\d{3}_")
ANOMALY_PRIORITY = (
    "附件缺失",
    "日期冲突",
    "日期待核验",
    "附件重复或无法匹配",
    "跨日期多订单",
)


@dataclass(frozen=True, slots=True)
class OrderDateInfo:
    source_file: str
    order_date: str | None
    source: str
    status: str
    evidence: str


def ensure_outside_repository(path: Path, label: str) -> Path:
    resolved = path.resolve()
    try:
        resolved.relative_to(PROJECT_ROOT)
    except ValueError:
        return resolved
    raise ValueError(f"{label}必须位于 Git 仓库外: {resolved}")


def sha256_path(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def split_filenames(value: object) -> list[str]:
    return [part.strip() for part in str(value or "").split(",") if part.strip()]


def load_master_records(path: Path) -> list[dict[str, str]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["发票记录"] if "发票记录" in workbook.sheetnames else workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    headers = [str(value or "").replace("\ufeff", "").strip() for value in next(rows)]
    records = []
    for row_number, row in enumerate(rows, start=2):
        record = {header: str(row[index] or "").strip() for index, header in enumerate(headers)}
        if not record.get("编号"):
            continue
        record["_row_number"] = str(row_number)
        record["上传人"] = normalize_uploader(record.get("上传人"))
        records.append(record)
    return records


def load_manual_dates(path: Path) -> dict[str, tuple[str | None, str]]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheet = workbook["黄金样本标注"] if "黄金样本标注" in workbook.sheetnames else workbook.worksheets[0]
    headers = [str(cell.value or "").strip() for cell in sheet[3]]
    source_index = headers.index("编号样本文件")
    date_index = headers.index("人工下单时间")
    result: dict[str, tuple[str | None, str]] = {}
    for row in sheet.iter_rows(min_row=4, values_only=True):
        source_name = MANUAL_PREFIX_RE.sub("", str(row[source_index] or "").strip())
        if not source_name:
            continue
        raw_date = row[date_index]
        if isinstance(raw_date, datetime):
            value = raw_date.date().isoformat()
        elif isinstance(raw_date, date):
            value = raw_date.isoformat()
        else:
            value = None
        result[source_name.casefold()] = (value, str(raw_date or "附件内未显示"))
    return result


def load_order_dates(rules_json: Path, manual_workbook: Path) -> dict[str, OrderDateInfo]:
    payload = json.loads(rules_json.read_text(encoding="utf-8"))
    manual = load_manual_dates(manual_workbook)
    result: dict[str, OrderDateInfo] = {}
    for item in payload.get("files", []):
        source_file = str(item["source_file"])
        manual_value = manual.get(source_file.casefold())
        if manual_value is not None:
            order_date, raw = manual_value
            result[source_file.casefold()] = OrderDateInfo(
                source_file,
                order_date,
                "人工标注",
                "人工确认" if order_date else "人工确认无日期",
                f"人工下单时间: {raw}",
            )
            continue
        raw_datetime = item.get("order_datetime")
        result[source_file.casefold()] = OrderDateInfo(
            source_file,
            str(raw_datetime)[:10] if raw_datetime else None,
            "最新版规则",
            str(item.get("status") or "NO_ORDER_DATE"),
            str(item.get("evidence_text") or ""),
        )
    return result


def _date_label(value: str | None) -> str:
    return value or "日期待核验"


def _amount_number(value: str) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def load_manual_reviews(path: Path | None) -> dict[str, dict[str, object]]:
    if path is None:
        return {}
    payload = json.loads(path.read_text(encoding="utf-8"))
    return {
        str(record.get("bx_id") or "").strip().casefold(): record
        for record in payload.get("records", [])
        if str(record.get("bx_id") or "").strip()
    }


def review_correction_files(review: dict[str, object] | None) -> list[Path]:
    if not review:
        return []
    directory = Path(str(review.get("correction_dir") or ""))
    return [directory / str(name) for name in review.get("correction_files", []) if (directory / str(name)).is_file()]


def attachment_type_for_path(path: Path) -> str:
    if "发票" in path.stem:
        return "发票"
    if "订单" in path.stem:
        return "订单"
    return "附件"


def review_is_special(review: dict[str, object] | None) -> bool:
    if not review:
        return False
    if str(review.get("review_status") or "") == "作废排除":
        return True
    notes = str(review.get("notes") or "")
    return any(keyword in notes for keyword in SPECIAL_REVIEW_KEYWORDS)


def build_plan(
    master_workbook: Path,
    attach_root: Path,
    rules_json: Path,
    manual_workbook: Path,
    reference_dir: Path,
    output_dir: Path,
    name_template: str = CLASSIFICATION_NAME_TEMPLATE,
    exclude_ids: set[str] | None = None,
    manual_review_json: Path | None = None,
    invoice_only_threshold: float = LOW_VALUE_INVOICE_ONLY_LIMIT,
) -> dict[str, object]:
    all_records = load_master_records(master_workbook)
    normalized_exclude_ids = {value.strip().casefold() for value in (exclude_ids or set()) if value.strip()}
    records = [record for record in all_records if record["编号"].casefold() not in normalized_exclude_ids]
    excluded_records = [record["编号"] for record in all_records if record["编号"].casefold() in normalized_exclude_ids]
    date_by_file = load_order_dates(rules_json, manual_workbook)
    manual_reviews = load_manual_reviews(manual_review_json)
    invoice_dir = attach_root / "发票"
    order_dir = attach_root / "订单截图"
    copies: list[dict[str, object]] = []
    anomalies: list[dict[str, object]] = []
    record_results: list[dict[str, object]] = []
    manual_review_results: list[dict[str, object]] = []
    source_hashes: dict[str, str] = {}
    target_sources: dict[str, str] = {}

    def add_copy(
        source: Path,
        target: Path,
        dimension: str,
        record: dict[str, str],
        attachment_type: str,
        order_date: str | None,
        date_source: str,
        status: str,
        anomaly: bool,
    ) -> None:
        source_key = str(source.resolve())
        target_key = str(target.resolve()).casefold()
        previous = target_sources.get(target_key)
        if previous == source_key:
            return
        if previous is not None:
            raise ValueError(f"分类目标重名: {target}")
        target_sources[target_key] = source_key
        if source_key not in source_hashes:
            source_hashes[source_key] = sha256_path(source)
        copies.append(
            {
                "dimension": dimension,
                "bx_id": record["编号"],
                "uploader": record["上传人"],
                "purpose": record.get("用途（必填）", ""),
                "amount": _amount_number(record.get("税后金额", "")),
                "attachment_type": attachment_type,
                "order_date": order_date,
                "date_source": date_source,
                "status": status,
                "is_anomaly": anomaly,
                "source_path": source_key,
                "target_path": str(target.resolve()),
                "target_relative": target.resolve().relative_to(output_dir.resolve()).as_posix(),
            }
        )

    for record in records:
        bx_id = record["编号"]
        uploader = record["上传人"]
        purpose = record.get("用途（必填）", "")
        amount = record.get("税后金额", "") or "未知金额"
        amount_value = _amount_number(amount)
        invoice_only_exempt = 0 < amount_value <= invoice_only_threshold
        order_required = not invoice_only_exempt
        review = manual_reviews.get(bx_id.casefold())
        review_status = str(review.get("review_status") or "") if review else ""
        review_date = str(review.get("confirmed_order_date") or review.get("current_order_date") or "") if review else ""
        review_date = review_date[:10] if review_date else None
        special_review = review_is_special(review)
        invoice_names = split_filenames(record.get("发票"))
        order_names = split_filenames(record.get("订单截图"))

        root_invoices = find_files_by_bx_id(invoice_dir, bx_id)
        override_invoices = [reference_dir / name for name in invoice_names if (reference_dir / name).is_file()]
        if override_invoices:
            selected_invoices = override_invoices
            extra_invoices: list[Path] = []
            superseded_invoices = root_invoices
        else:
            expected_invoice_count = count_comma_segments(record.get("发票", ""))
            selected_invoices = root_invoices[:expected_invoice_count] if expected_invoice_count else []
            extra_invoices = root_invoices[expected_invoice_count:] if expected_invoice_count else root_invoices
            superseded_invoices = []

        expected_order_count = count_comma_segments(record.get("订单截图", ""))
        root_orders = find_files_by_bx_id(order_dir, bx_id)
        selected_orders = root_orders[:expected_order_count] if expected_order_count else []
        extra_orders = root_orders[expected_order_count:] if expected_order_count else root_orders

        correction_files = review_correction_files(review)
        correction_invoices = [path for path in correction_files if attachment_type_for_path(path) == "发票"]
        correction_orders = [path for path in correction_files if attachment_type_for_path(path) == "订单"]
        if review_status in RESOLVED_REVIEW_STATUSES and correction_files:
            if correction_invoices:
                selected_invoices = correction_invoices
            if correction_orders:
                selected_orders = correction_orders
            extra_invoices = []
            extra_orders = []

        if review_status == "作废排除":
            special_sources = correction_files or (selected_invoices + extra_invoices + selected_orders + extra_orders)
            totals = {
                kind: sum(1 for path in special_sources if attachment_type_for_path(path) == kind)
                for kind in {"发票", "订单", "附件"}
            }
            counters = {"发票": 0, "订单": 0, "附件": 0}
            for source in special_sources:
                attachment_type = attachment_type_for_path(source)
                counters[attachment_type] += 1
                filename = render_target_filename(
                    name_template, bx_id, uploader, purpose, amount, attachment_type,
                    counters[attachment_type], max(1, totals[attachment_type]), source.suffix,
                    order_date=review_date or "日期不适用",
                )
                add_copy(
                    source, output_dir / "04_作废或特殊处理" / "作废排除" / bx_id / f"作废排除_{filename}",
                    "特殊处理", record, attachment_type, review_date, "人工复核", "作废排除", False,
                )
            manual_review_results.append(
                {
                    **review,
                    "result": "已从正常日期和人员分类排除",
                    "special_flag": True,
                    "used_files": [path.name for path in special_sources],
                }
            )
            continue

        order_entries = []
        for path in selected_orders:
            if review_status in RESOLVED_REVIEW_STATUSES:
                info = OrderDateInfo(
                    path.name, review_date, "人工复核",
                    "人工确认" if review_date else "人工确认日期不适用",
                    str(review.get("notes") or "人工复核确认"),
                )
            else:
                info = date_by_file.get(path.name.casefold()) or OrderDateInfo(
                    path.name, None, "无识别结果", "NO_ORDER_DATE", "订单附件未出现在规则结果中"
                )
            order_entries.append((path, info))

        dates = ([review_date] if review_status in RESOLVED_REVIEW_STATUSES and review_date else
                 sorted({info.order_date for _, info in order_entries if info.order_date}))
        anomaly_types: list[str] = []
        evidence: list[str] = []
        missing: list[str] = []
        if not selected_invoices:
            anomaly_types.append("附件缺失")
            missing.append("发票")
        if order_required and not selected_orders:
            anomaly_types.append("附件缺失")
            missing.append("订单")
        if extra_invoices or (order_required and extra_orders):
            anomaly_types.append("附件重复或无法匹配")
            evidence.append(
                f"多余发票{len(extra_invoices)}个，多余订单{len(extra_orders) if order_required else 0}个"
            )
        if order_required and len(dates) > 1:
            anomaly_types.append("跨日期多订单")
            evidence.append("同一编号包含多个下单日期: " + ", ".join(dates))
        for _, info in order_entries:
            evidence.append(f"{info.source_file}: {info.source}/{info.status}; {info.evidence}")
            if invoice_only_exempt:
                continue
            if info.status == "CONFLICT":
                anomaly_types.append("日期冲突")
            elif info.status == "DUPLICATE":
                anomaly_types.append("附件重复或无法匹配")
            elif info.status not in {"AUTO_ACCEPTED", "人工确认"}:
                anomaly_types.append("日期待核验")
        anomaly_types = [value for value in ANOMALY_PRIORITY if value in set(anomaly_types)]
        if review_status in RESOLVED_REVIEW_STATUSES:
            anomaly_types = []
            missing = []
            evidence.append(f"人工复核: {review_status}; {str(review.get('notes') or '')}")
        is_anomaly = bool(anomaly_types)
        undated_label = ("日期不适用" if review_status in RESOLVED_REVIEW_STATUSES else
                         ("低额免订单截图" if invoice_only_exempt else "日期待核验"))

        def apply_special_prefix(filename: str) -> str:
            return f"特殊处理_{filename}" if special_review else filename

        # 人员分类：订单按自身日期命名；发票在多日期记录中按日期各复制一份。
        for index, (source, info) in enumerate(order_entries, start=1):
            filename = render_target_filename(
                name_template, bx_id, uploader, purpose, amount, "订单", index, max(1, len(order_entries)),
                source.suffix, order_date=info.order_date or undated_label,
            )
            filename = apply_special_prefix(filename)
            add_copy(source, output_dir / "02_按人员分类" / safe_text(uploader) / filename, "人员", record,
                     "订单", info.order_date, info.source, info.status, is_anomaly)
        invoice_dates = dates or [None]
        for order_date in invoice_dates:
            for index, source in enumerate(selected_invoices, start=1):
                filename = render_target_filename(
                    name_template, bx_id, uploader, purpose, amount, "发票", index, max(1, len(selected_invoices)),
                    source.suffix, order_date=order_date or undated_label,
                )
                invoice_date_source = "记录级日期" if order_date else (
                    "低额免订单截图" if invoice_only_exempt else "记录级日期"
                )
                invoice_status = "已匹配" if order_date else (
                    "免订单" if invoice_only_exempt else "已匹配"
                )
                filename = apply_special_prefix(filename)
                add_copy(source, output_dir / "02_按人员分类" / safe_text(uploader) / filename, "人员", record,
                         "发票", order_date, invoice_date_source, invoice_status, is_anomaly)

        # 日期分类：无日期附件不进入日期目录；低置信但有日期的记录仍双放。
        for index, (source, info) in enumerate(order_entries, start=1):
            if not info.order_date:
                continue
            filename = render_target_filename(
                name_template, bx_id, uploader, purpose, amount, "订单", index, max(1, len(order_entries)),
                source.suffix, order_date=info.order_date,
            )
            filename = apply_special_prefix(filename)
            add_copy(source, output_dir / "01_按下单日期分类" / info.order_date / filename, "日期", record,
                     "订单", info.order_date, info.source, info.status, is_anomaly)
        for order_date in dates:
            for index, source in enumerate(selected_invoices, start=1):
                filename = render_target_filename(
                    name_template, bx_id, uploader, purpose, amount, "发票", index, max(1, len(selected_invoices)),
                    source.suffix, order_date=order_date,
                )
                filename = apply_special_prefix(filename)
                add_copy(source, output_dir / "01_按下单日期分类" / order_date / filename, "日期", record,
                         "发票", order_date, "记录级日期", "已匹配", is_anomaly)

        if is_anomaly:
            category = anomaly_types[0]
            anomaly_root = output_dir / "03_异常待复核" / category / bx_id
            anomaly_date = dates[0] if len(dates) == 1 else ("多日期" if dates else None)
            all_invoices = selected_invoices + extra_invoices
            all_orders = selected_orders + extra_orders
            for index, source in enumerate(all_invoices, start=1):
                filename = render_target_filename(
                    name_template, bx_id, uploader, purpose, amount, "发票", index, max(1, len(all_invoices)),
                    source.suffix, order_date=_date_label(anomaly_date),
                )
                add_copy(source, anomaly_root / filename, "异常", record, "发票", anomaly_date,
                         "异常汇总", ";".join(anomaly_types), True)
            for index, source in enumerate(all_orders, start=1):
                info = date_by_file.get(source.name.casefold())
                filename = render_target_filename(
                    name_template, bx_id, uploader, purpose, amount, "订单", index, max(1, len(all_orders)),
                    source.suffix, order_date=_date_label(info.order_date if info else anomaly_date),
                )
                add_copy(source, anomaly_root / filename, "异常", record, "订单",
                         info.order_date if info else anomaly_date, info.source if info else "无识别结果",
                         info.status if info else "NO_ORDER_DATE", True)
            anomalies.append(
                {
                    "bx_id": bx_id,
                    "uploader": uploader,
                    "purpose": purpose,
                    "amount": _amount_number(amount),
                    "dates": dates,
                    "types": anomaly_types,
                    "primary_category": category,
                    "evidence": " | ".join(dict.fromkeys(evidence)),
                    "missing": missing,
                    "available_invoices": [path.name for path in selected_invoices + extra_invoices],
                    "available_orders": [path.name for path in selected_orders + extra_orders],
                    "suggestion": "核对下单日期及订单、发票附件数量；确认后更新主表或人工标注。",
                }
            )

        if review_status in RESOLVED_REVIEW_STATUSES:
            used_files = selected_invoices + selected_orders
            if special_review:
                totals = {"发票": len(selected_invoices), "订单": len(selected_orders)}
                for attachment_type, sources in (("发票", selected_invoices), ("订单", selected_orders)):
                    for index, source in enumerate(sources, start=1):
                        filename = render_target_filename(
                            name_template, bx_id, uploader, purpose, amount, attachment_type, index,
                            max(1, totals[attachment_type]), source.suffix,
                            order_date=review_date or "日期不适用",
                        )
                        add_copy(
                            source,
                            output_dir / "04_作废或特殊处理" / "特殊标记" / bx_id / f"特殊处理_{filename}",
                            "特殊处理", record, attachment_type, review_date, "人工复核", review_status, False,
                        )
            manual_review_results.append(
                {
                    **review,
                    "result": "已按人工结论重新分类",
                    "special_flag": special_review,
                    "used_files": [path.name for path in used_files],
                }
            )

        record_results.append(
            {
                "bx_id": bx_id,
                "uploader": uploader,
                "purpose": purpose,
                "amount": _amount_number(amount),
                "dates": dates,
                "is_anomaly": is_anomaly,
                "anomaly_types": anomaly_types,
                "invoice_count": len(selected_invoices),
                "order_count": len(selected_orders),
                "invoice_only_exempt": invoice_only_exempt,
                "order_required": order_required,
                "manual_review_status": review_status,
                "special_flag": special_review,
                "superseded_invoices": [path.name for path in superseded_invoices],
            }
        )

    date_groups = sorted({entry["order_date"] for entry in copies if entry["dimension"] == "日期"})
    people = sorted({entry["uploader"] for entry in copies if entry["dimension"] == "人员"})
    return {
        "version": 1,
        "generated_at": datetime.now().astimezone().isoformat(timespec="seconds"),
        "output_root": str(output_dir.resolve()),
        "inputs": {
            "master_workbook": str(master_workbook.resolve()),
            "attach_root": str(attach_root.resolve()),
            "rules_json": str(rules_json.resolve()),
            "manual_workbook": str(manual_workbook.resolve()),
            "reference_dir": str(reference_dir.resolve()),
            "excluded_ids": sorted(excluded_records),
            "manual_review_json": str(manual_review_json.resolve()) if manual_review_json else None,
        },
        "stats": {
            "records": len(record_results),
            "excluded_records": len(excluded_records) + sum(
                1 for review in manual_review_results if review.get("review_status") == "作废排除"
            ),
            "manual_review_records": len(manual_review_results),
            "manual_excluded_records": sum(
                1 for review in manual_review_results if review.get("review_status") == "作废排除"
            ),
            "special_records": sum(1 for review in manual_review_results if review.get("special_flag")),
            "copies": len(copies),
            "date_groups": len(date_groups),
            "people": len(people),
            "anomaly_records": len(anomalies),
            "invoice_only_records": sum(1 for record in record_results if record["invoice_only_exempt"]),
            "source_files": len(source_hashes),
        },
        "records": record_results,
        "copies": copies,
        "anomalies": anomalies,
        "manual_reviews": manual_review_results,
        "source_hashes": source_hashes,
    }


def execute_plan(plan: dict[str, object], dry_run: bool) -> None:
    if dry_run:
        return
    output_root = Path(str(plan["output_root"]))
    if output_root.exists() and any(output_root.iterdir()):
        raise FileExistsError(f"输出目录已存在且非空，请更换目录: {output_root}")
    output_root.mkdir(parents=True, exist_ok=True)
    for entry in plan["copies"]:
        source = Path(str(entry["source_path"]))
        target = Path(str(entry["target_path"]))
        target.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(source, target)
    manifest_path = output_root / "classification_manifest.json"
    manifest_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding="utf-8")


def parse_args(argv: list[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="按下单日期和上传人分类发票、订单附件")
    parser.add_argument("--workbook", required=True, type=Path, help="整理后的主工作簿")
    parser.add_argument("--attach-root", required=True, type=Path, help="包含发票和订单截图目录的附件根目录")
    parser.add_argument("--rules-json", required=True, type=Path, help="最新版订单日期规则结果 JSON")
    parser.add_argument("--manual-labels", required=True, type=Path, help="人工标注工作簿")
    parser.add_argument("--reference-dir", required=True, type=Path, help="更新版发票文件所在目录")
    parser.add_argument("--output-dir", required=True, type=Path)
    parser.add_argument("--name-template", default=CLASSIFICATION_NAME_TEMPLATE)
    parser.add_argument("--exclude-id", action="append", default=[], help="从分类和报表中排除的编号，可重复使用")
    parser.add_argument("--manual-review-json", type=Path, help="由人工复核清单解析得到的 JSON，人工结论优先")
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument(
        "--invoice-only-threshold",
        type=float,
        default=LOW_VALUE_INVOICE_ONLY_LIMIT,
        help="仅要求发票、不要求订单截图的最高金额，默认 200",
    )
    return parser.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv)
    if args.invoice_only_threshold < 0:
        raise ValueError("--invoice-only-threshold 不能小于 0")
    output_dir = ensure_outside_repository(args.output_dir, "分类结果")
    for source in (args.workbook, args.rules_json, args.manual_labels, args.manual_review_json):
        if source is None:
            continue
        if not source.is_file():
            raise FileNotFoundError(source)
    if not (args.attach_root / "发票").is_dir() or not (args.attach_root / "订单截图").is_dir():
        raise NotADirectoryError("附件根目录必须包含“发票”和“订单截图”")
    plan = build_plan(
        args.workbook.resolve(), args.attach_root.resolve(), args.rules_json.resolve(),
        args.manual_labels.resolve(), args.reference_dir.resolve(), output_dir, args.name_template,
        exclude_ids=set(args.exclude_id), manual_review_json=args.manual_review_json.resolve() if args.manual_review_json else None,
        invoice_only_threshold=args.invoice_only_threshold,
    )
    execute_plan(plan, args.dry_run)
    print(json.dumps(plan["stats"], ensure_ascii=False, indent=2))
    if args.dry_run:
        print("预演完成：未复制文件，也未写入清单。")
    else:
        print(f"分类完成: {output_dir}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
