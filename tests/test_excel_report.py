from datetime import datetime

from openpyxl import load_workbook

from order_date.excel_report import ReportRunStats, write_excel_report
from order_date.models import (
    AggregatedResult,
    ExpenseRecord,
    ExtractionResult,
    OCRPage,
    OCRTextBlock,
    SourceFile,
)


def test_excel_report_has_four_sheets_typed_dates_and_links(tmp_path):
    attachment = tmp_path / "BX123.jpg"
    attachment.write_bytes(b"sample")
    source = SourceFile("BX123", attachment, attachment.name, "hash", ".jpg", 6, 1)
    moment = datetime(2026, 8, 16, 12, 34, 56)
    result = ExtractionResult(
        "BX123",
        attachment.name,
        "hash",
        "jd",
        20,
        "12345678901234567890",
        moment,
        80,
        "NEEDS_REVIEW",
        "标签 下单时间",
        "rules-v1",
        1,
    )
    aggregated = AggregatedResult(
        "BX123",
        result.order_number,
        moment,
        80,
        "NEEDS_REVIEW",
        (attachment.name,),
        result.evidence_text,
        result.rule_version,
    )
    block = OCRTextBlock(
        "2026-08-16 12:34:56",
        "2026-08-16 12:34:56",
        0.99,
        ((0, 0), (100, 0), (100, 20), (0, 20)),
        0,
    )
    output = write_excel_report(
        tmp_path / "report.xlsx",
        (ExpenseRecord("BX123", "测试人", "测试用途", "12.34", "", attachment.name, 2),),
        (source,),
        (result,),
        (aggregated,),
        {"hash": (OCRPage(0, (block,)),)},
        frozenset({"hash"}),
        {attachment.name: 0.123},
        ReportRunStats(1, 1, 1.5, "engine-v1"),
    )

    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["订单日期汇总", "附件识别明细", "待人工复核", "运行统计"]
    summary = workbook["订单日期汇总"]
    assert isinstance(summary["G2"].value, datetime)
    assert isinstance(summary["H2"].value, datetime)
    assert summary["F2"].value == "ID 12345678901234567890"
    assert summary["F2"].number_format == "@"
    assert summary["L2"].hyperlink is not None
    assert summary["L2"].hyperlink.target == attachment.as_uri()
    assert summary["I2"].value == "待复核"
    review = workbook["待人工复核"]
    assert review.max_row == 2
    assert review["B2"].hyperlink is not None
    assert "平台" not in review["E2"].value
    assert review["F2"].number_format == "yyyy-mm-dd"
    assert workbook["附件识别明细"].max_row == 2
    assert workbook["运行统计"]["A2"].value == "总文件数"
    workbook.close()
