from datetime import date
from decimal import Decimal

from openpyxl import load_workbook

from invoice_issuer.excel_report import write_invoice_issuer_report
from invoice_issuer.models import DailyMerchantGroup, InvoiceIssuerResult, ReviewItem


def test_writes_self_contained_invoice_issuer_workbook(tmp_path):
    group = DailyMerchantGroup(
        date(2026, 1, 2), "示例公司", "示例公司", Decimal("1200"), "EXCEEDED",
        ("DEMO001",), ("DEMO001.pdf",), "合计 1200 元",
    )
    result = InvoiceIssuerResult(
        "DEMO001", "DEMO001.pdf", "hash", "示例公司", "示例公司", 90,
        "AUTO_ACCEPTED", "销售方名称", 1, 0,
    )
    review = ReviewItem("DEMO002", ("DEMO002.pdf",), "缺少唯一订单日期")
    output = tmp_path / "report.xlsx"
    write_invoice_issuer_report(
        output,
        groups=(group,),
        results=(result,),
        review_items=(review,),
        daily_limit=Decimal("1000"),
        stats={"发票文件数": 1},
    )
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["单日同公司汇总", "发票识别明细", "待人工复核", "运行统计"]
    assert workbook["单日同公司汇总"]["D2"].value == "超额"
    assert workbook["运行统计"]["B2"].value == 1000
