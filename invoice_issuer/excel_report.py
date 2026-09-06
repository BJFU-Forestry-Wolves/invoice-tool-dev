from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

from .models import DailyMerchantGroup, InvoiceIssuerResult, ReviewItem


HEADER_FILL = PatternFill("solid", fgColor="2563EB")
HEADER_FONT = Font(color="FFFFFF", bold=True)
REVIEW_FILL = PatternFill("solid", fgColor="FFF7D6")
ERROR_FILL = PatternFill("solid", fgColor="FDECEC")


def _finish_sheet(sheet, widths: tuple[int, ...], table_name: str) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    if sheet.max_row >= 2:
        table = Table(displayName=table_name, ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        sheet.add_table(table)


def write_invoice_issuer_report(
    output_path: Path,
    *,
    groups: tuple[DailyMerchantGroup, ...],
    results: tuple[InvoiceIssuerResult, ...],
    review_items: tuple[ReviewItem, ...],
    daily_limit: Decimal,
    stats: dict[str, object],
) -> Path:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "单日同公司汇总"
    summary.append(["订单日期", "开票公司", "当日合计（元）", "核验状态", "报销编号", "发票文件", "证据"])
    for item in groups:
        summary.append([
            item.order_date,
            item.issuer_name,
            float(item.total_amount),
            "超额" if item.status == "EXCEEDED" else "通过",
            "、".join(item.bx_ids),
            "、".join(item.source_files),
            item.evidence_text,
        ])
        summary.cell(summary.max_row, 1).number_format = "yyyy-mm-dd"
        summary.cell(summary.max_row, 3).number_format = "#,##0.00"
        if item.status == "EXCEEDED":
            for cell in summary[summary.max_row]:
                cell.fill = ERROR_FILL
    _finish_sheet(summary, (13, 34, 16, 12, 28, 36, 48), "DailyMerchantTable")

    details = workbook.create_sheet("发票识别明细")
    details.append(["报销编号", "发票文件", "销售方名称", "规范化名称", "置信度", "识别状态", "页码", "候选数", "证据", "错误代码"])
    for item in results:
        details.append([
            item.bx_id,
            item.source_file,
            item.issuer_name or "",
            item.normalized_issuer or "",
            item.confidence_score / 100,
            item.status,
            item.page_index + 1 if item.page_index is not None else None,
            item.candidate_count,
            item.evidence_text,
            item.error_code or "",
        ])
        details.cell(details.max_row, 1).number_format = "@"
        details.cell(details.max_row, 5).number_format = "0%"
    _finish_sheet(details, (18, 28, 34, 34, 12, 18, 10, 10, 52, 20), "InvoiceDetailTable")

    review = workbook.create_sheet("待人工复核")
    review.append(["报销编号", "发票文件", "复核原因", "OCR销售方", "关联日期", "报销金额", "人工确认销售方", "人工确认日期", "复核人", "备注"])
    for item in review_items:
        review.append([
            item.bx_id,
            "、".join(item.source_files),
            item.reason,
            "、".join(item.issuer_names),
            "、".join(value.isoformat() for value in item.order_dates),
            float(item.amount) if item.amount is not None else None,
            "",
            None,
            "",
            "",
        ])
        review.cell(review.max_row, 1).number_format = "@"
        review.cell(review.max_row, 6).number_format = "#,##0.00"
        review.cell(review.max_row, 8).number_format = "yyyy-mm-dd"
        for cell in review[review.max_row][6:10]:
            cell.fill = REVIEW_FILL
    _finish_sheet(review, (18, 34, 38, 32, 24, 15, 28, 18, 16, 34), "IssuerReviewTable")

    run_stats = workbook.create_sheet("运行统计")
    run_stats.append(["指标", "值"])
    run_stats.append(["单日同公司限额（元）", float(daily_limit)])
    for key, value in stats.items():
        run_stats.append([key, value])
    run_stats.append(["报告生成时间", datetime.now().isoformat(timespec="seconds")])
    run_stats.cell(2, 2).number_format = "#,##0.00"
    _finish_sheet(run_stats, (32, 64), "IssuerStatsTable")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path
