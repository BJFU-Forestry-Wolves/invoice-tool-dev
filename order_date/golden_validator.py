from __future__ import annotations

import json
import re
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import load_workbook


NUMBERED_PREFIX_RE = re.compile(r"^\d{3}_")


@dataclass(frozen=True, slots=True)
class DateValidationRow:
    row_number: int
    source_file: str
    status: str
    predicted_date: date | None
    manual_date: date | None

    @property
    def correct(self) -> bool:
        return self.predicted_date == self.manual_date


@dataclass(frozen=True, slots=True)
class DateValidationSummary:
    rows: tuple[DateValidationRow, ...]

    @property
    def total(self) -> int:
        return len(self.rows)

    @property
    def correct(self) -> int:
        return sum(row.correct for row in self.rows)

    @property
    def dated_rows(self) -> tuple[DateValidationRow, ...]:
        return tuple(row for row in self.rows if row.manual_date is not None)

    @property
    def no_date_rows(self) -> tuple[DateValidationRow, ...]:
        return tuple(row for row in self.rows if row.manual_date is None)

    @property
    def auto_accepted_rows(self) -> tuple[DateValidationRow, ...]:
        return tuple(row for row in self.rows if row.status == "AUTO_ACCEPTED")

    @staticmethod
    def accuracy(rows: tuple[DateValidationRow, ...]) -> float:
        return sum(row.correct for row in rows) / len(rows) if rows else 0.0


def _source_basename(value: str) -> str:
    return NUMBERED_PREFIX_RE.sub("", value.replace("\\", "/").rsplit("/", 1)[-1])


def _as_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip().replace("/", "-").replace(".", "-")
    match = re.search(r"20\d{2}-\d{1,2}-\d{1,2}", text)
    if match:
        return datetime.fromisoformat(match.group(0)).date()
    if any(character.isdigit() for character in text):
        raise ValueError(f"无法解析人工日期: {value}")
    return None


def validate_golden_dates(prediction_json: Path, workbook_path: Path) -> DateValidationSummary:
    payload = json.loads(prediction_json.read_text(encoding="utf-8"))
    predictions = {}
    for item in payload["files"]:
        key = _source_basename(item["source_file"])
        if key in predictions:
            raise ValueError(f"预测结果存在重复文件名: {key}")
        predictions[key] = item

    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    try:
        sheet = workbook["黄金样本标注"]
        header_row = None
        headers = {}
        for candidate_row in range(1, min(sheet.max_row, 20) + 1):
            candidate_headers = {
                sheet.cell(candidate_row, column).value: column for column in range(1, sheet.max_column + 1)
            }
            if ("编号样本文件" in candidate_headers or "源文件" in candidate_headers) and "人工下单时间" in candidate_headers:
                header_row = candidate_row
                headers = candidate_headers
                break
        if header_row is None:
            raise ValueError("黄金样本缺少表头")
        source_column = headers.get("编号样本文件") or headers.get("源文件")
        manual_column = headers.get("人工下单时间")
        if source_column is None or manual_column is None:
            raise ValueError("黄金样本缺少源文件或人工下单时间列")

        rows = []
        for row_number in range(header_row + 1, sheet.max_row + 1):
            source_value = sheet.cell(row_number, source_column).value
            if not source_value:
                continue
            source_file = _source_basename(str(source_value))
            prediction = predictions.get(source_file)
            if prediction is None:
                raise ValueError(f"找不到预测结果: {source_file}")
            predicted_date = _as_date(prediction.get("order_datetime"))
            manual_date = _as_date(sheet.cell(row_number, manual_column).value)
            rows.append(
                DateValidationRow(
                    row_number=row_number,
                    source_file=source_file,
                    status=prediction["status"],
                    predicted_date=predicted_date,
                    manual_date=manual_date,
                )
            )
    finally:
        workbook.close()
    return DateValidationSummary(tuple(rows))
