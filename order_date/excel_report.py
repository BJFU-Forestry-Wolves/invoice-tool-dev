from __future__ import annotations

from collections import Counter
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Mapping

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .candidate_extractor import extract_date_candidates
from .models import AggregatedResult, ExpenseRecord, ExtractionResult, OCRPage, SourceFile


STATUS_LABELS = {
    "AUTO_ACCEPTED": "自动通过",
    "NEEDS_REVIEW": "待复核",
    "NO_ORDER_DATE": "未确认下单日期",
    "OCR_FAILED": "OCR失败",
    "CONFLICT": "冲突",
    "DUPLICATE": "重复附件",
}
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
AUTO_FILL = PatternFill("solid", fgColor="E2F0D9")
REVIEW_FILL = PatternFill("solid", fgColor="FFF2CC")
ERROR_FILL = PatternFill("solid", fgColor="FCE4D6")
INPUT_FILL = PatternFill("solid", fgColor="D9EAF7")


@dataclass(frozen=True, slots=True)
class ReportRunStats:
    rule_cache_hits: int
    ocr_cache_hits: int
    elapsed_seconds: float
    engine_fingerprint: str
    scan_issue_count: int = 0


def _status_label(status: str) -> str:
    return STATUS_LABELS.get(status, status)


def _amount(value: str):
    text = value.strip().replace(",", "").replace("¥", "").replace("￥", "")
    try:
        return float(text)
    except ValueError:
        return value


def _add_table(sheet, name: str) -> None:
    if sheet.max_row < 2 or sheet.max_column < 1:
        return
    table = Table(displayName=name, ref=sheet.dimensions)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)


def _style_sheet(sheet, widths: tuple[int, ...]) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.sheet_view.showGridLines = False
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 30
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def _add_status_formatting(sheet, status_column: str, end_row: int) -> None:
    if end_row < 2:
        return
    target = f"{status_column}2:{status_column}{end_row}"
    sheet.conditional_formatting.add(
        target,
        FormulaRule(formula=[f'${status_column}2="自动通过"'], fill=AUTO_FILL),
    )
    sheet.conditional_formatting.add(
        target,
        FormulaRule(
            formula=[f'OR(${status_column}2="待复核",${status_column}2="未确认下单日期")'],
            fill=REVIEW_FILL,
        ),
    )
    sheet.conditional_formatting.add(
        target,
        FormulaRule(formula=[f'OR(${status_column}2="冲突",${status_column}2="OCR失败")'], fill=ERROR_FILL),
    )


def _source_link(cell, source: SourceFile | None) -> None:
    if source is None:
        return
    cell.value = source.relative_path
    cell.hyperlink = source.path.as_uri()
    cell.style = "Hyperlink"


def _suggestion(status: str) -> str:
    if status == "CONFLICT":
        return "核对多个日期候选及同一报销编号下的不同附件"
    if status == "OCR_FAILED":
        return "检查文件能否打开，必要时重新截图或转为 PNG/PDF"
    if status == "NO_ORDER_DATE":
        return "确认页面是否展示下单/创建日期；若只有付款或成交日期请备注"
    return "核对日期标签和候选日期后填写人工确认日期"


def write_excel_report(
    output_path: Path,
    records: tuple[ExpenseRecord, ...],
    sources: tuple[SourceFile, ...],
    results: tuple[ExtractionResult, ...],
    aggregated: tuple[AggregatedResult, ...],
    pages_by_hash: Mapping[str, tuple[OCRPage, ...]],
    ocr_cache_hit_hashes: frozenset[str],
    processing_seconds: Mapping[str, float],
    stats: ReportRunStats,
) -> Path:
    record_by_bx = {record.bx_id.upper(): record for record in records}
    source_by_name = {source.relative_path: source for source in sources}
    result_by_name = {result.source_file: result for result in results}

    workbook = Workbook()
    summary = workbook.active
    summary.title = "订单日期汇总"
    detail = workbook.create_sheet("附件识别明细")
    review = workbook.create_sheet("待人工复核")
    run_stats = workbook.create_sheet("运行统计")

    summary.append(
        [
            "报销编号", "上传人", "用途", "税后金额", "电商平台", "订单号", "下单日期", "下单时间",
            "识别状态", "置信度", "证据文字", "证据文件", "识别方式", "规则版本",
        ]
    )
    for item in aggregated:
        record = record_by_bx.get(item.bx_id.upper())
        first_source_name = item.source_files[0] if item.source_files else None
        source = source_by_name.get(first_source_name) if first_source_name else None
        related = [result_by_name[name] for name in item.source_files if name in result_by_name]
        platform = related[0].platform if related else "unknown"
        cache_only = bool(related) and all(result.file_hash in ocr_cache_hit_hashes for result in related)
        summary.append(
            [
                item.bx_id,
                record.uploader if record else "",
                record.purpose if record else "",
                _amount(record.amount) if record else "",
                platform,
                f"ID {item.order_number}" if item.order_number else "",
                item.order_datetime.date() if item.order_datetime else None,
                item.order_datetime,
                _status_label(item.status),
                item.confidence_score,
                item.evidence_text,
                "",
                "PaddleOCR（缓存）" if cache_only else "PaddleOCR",
                item.rule_version,
            ]
        )
        row = summary.max_row
        _source_link(summary.cell(row, 12), source)
        summary.cell(row, 1).number_format = "@"
        summary.cell(row, 6).number_format = "@"
        summary.cell(row, 7).number_format = "yyyy-mm-dd"
        summary.cell(row, 8).number_format = "yyyy-mm-dd hh:mm:ss"
        summary.cell(row, 10).number_format = "0"
        if isinstance(summary.cell(row, 4).value, float):
            summary.cell(row, 4).number_format = "#,##0.00"

    detail.append(
        [
            "报销编号", "源文件", "文件哈希", "页码", "平台", "平台得分", "全部日期候选", "最终选择",
            "识别状态", "OCR最低置信度", "OCR平均置信度", "规则处理耗时(秒)", "缓存命中", "异常码",
        ]
    )
    for result in results:
        pages = pages_by_hash.get(result.file_hash, ())
        source = source_by_name.get(result.source_file)
        if not pages:
            pages = (OCRPage(0, ()),)
        for page in pages:
            confidences = [block.confidence for block in page.blocks]
            candidates = extract_date_candidates((page,))
            candidate_text = "; ".join(candidate.normalized_datetime.isoformat(sep=" ") for candidate in candidates)
            detail.append(
                [
                    result.bx_id,
                    "",
                    result.file_hash,
                    page.page_index + 1,
                    result.platform,
                    result.platform_score,
                    candidate_text,
                    result.order_datetime,
                    _status_label(result.status),
                    min(confidences) if confidences else None,
                    sum(confidences) / len(confidences) if confidences else None,
                    processing_seconds.get(result.source_file),
                    "是" if result.file_hash in ocr_cache_hit_hashes else "否",
                    result.error_code or "",
                ]
            )
            row = detail.max_row
            _source_link(detail.cell(row, 2), source)
            detail.cell(row, 1).number_format = "@"
            detail.cell(row, 3).number_format = "@"
            detail.cell(row, 8).number_format = "yyyy-mm-dd hh:mm:ss"
            for column in (10, 11):
                detail.cell(row, column).number_format = "0.0%"
            detail.cell(row, 12).number_format = "0.000"

    review.append(
        [
            "报销编号", "源文件", "候选日期", "复核原因", "建议检查项", "人工确认日期", "复核人", "复核时间", "备注",
        ]
    )
    review_statuses = {"NEEDS_REVIEW", "NO_ORDER_DATE", "CONFLICT", "OCR_FAILED"}
    for result in results:
        if result.status not in review_statuses:
            continue
        source = source_by_name.get(result.source_file)
        review.append(
            [
                result.bx_id,
                "",
                result.order_datetime,
                f"{_status_label(result.status)}：{result.evidence_text}",
                _suggestion(result.status),
                None,
                "",
                None,
                "",
            ]
        )
        row = review.max_row
        _source_link(review.cell(row, 2), source)
        review.cell(row, 1).number_format = "@"
        review.cell(row, 3).number_format = "yyyy-mm-dd hh:mm:ss"
        review.cell(row, 6).number_format = "yyyy-mm-dd"
        review.cell(row, 8).number_format = "yyyy-mm-dd hh:mm:ss"
        for column in range(6, 10):
            review.cell(row, column).fill = INPUT_FILL

    result_statuses = Counter(result.status for result in results)
    platforms = Counter(result.platform for result in results)
    extensions = Counter(source.extension for source in sources)
    successful_extensions = Counter(
        source_by_name[result.source_file].extension
        for result in results
        if result.source_file in source_by_name and result.status != "OCR_FAILED"
    )
    total_pages = sum(len(pages_by_hash.get(result.file_hash, ())) for result in results)
    stats_rows = [
        ("总文件数", len(results)),
        ("总页数", total_pages),
        ("OCR缓存命中数", stats.ocr_cache_hits),
        ("规则缓存命中数", stats.rule_cache_hits),
        ("自动通过数", result_statuses["AUTO_ACCEPTED"]),
        ("待复核数", result_statuses["NEEDS_REVIEW"] + result_statuses["NO_ORDER_DATE"]),
        ("冲突数", result_statuses["CONFLICT"]),
        ("失败数", result_statuses["OCR_FAILED"]),
        ("重复数", result_statuses["DUPLICATE"]),
        ("扫描跳过数", stats.scan_issue_count),
        ("总耗时(秒)", round(stats.elapsed_seconds, 3)),
        ("规则版本", results[0].rule_version if results else ""),
        ("模型指纹", stats.engine_fingerprint),
    ]
    stats_rows.extend((f"平台:{platform}", count) for platform, count in sorted(platforms.items()))
    stats_rows.extend(
        (f"格式成功率:{extension}", successful_extensions[extension] / count if count else 0)
        for extension, count in sorted(extensions.items())
    )
    run_stats.append(["指标", "数值"])
    for label, value in stats_rows:
        run_stats.append([label, value])
        if label.startswith("格式成功率:"):
            run_stats.cell(run_stats.max_row, 2).number_format = "0.0%"

    _style_sheet(summary, (18, 14, 28, 14, 13, 28, 13, 21, 18, 10, 46, 28, 18, 22))
    _style_sheet(detail, (18, 28, 18, 8, 12, 11, 44, 21, 18, 14, 14, 16, 10, 20))
    _style_sheet(review, (18, 28, 21, 48, 42, 16, 14, 21, 30))
    _style_sheet(run_stats, (28, 36))
    for row in range(2, summary.max_row + 1):
        summary.row_dimensions[row].height = 48 if len(str(summary.cell(row, 11).value or "")) > 70 else 30
    for row in range(2, detail.max_row + 1):
        detail.row_dimensions[row].height = 58
    for row in range(2, review.max_row + 1):
        review.row_dimensions[row].height = 66
    _add_status_formatting(summary, "I", summary.max_row)
    _add_status_formatting(detail, "I", detail.max_row)
    _add_table(summary, "OrderDateSummaryTable")
    _add_table(detail, "AttachmentDetailTable")
    _add_table(run_stats, "RunStatisticsTable")

    workbook.properties.title = "订单下单日期识别结果"
    workbook.properties.created = datetime.now()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(output_path)
        return output_path
    except PermissionError:
        fallback = output_path.with_name(f"{output_path.stem}_{datetime.now():%Y%m%d_%H%M%S}{output_path.suffix}")
        workbook.save(fallback)
        return fallback
